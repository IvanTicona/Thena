"""
Tests for the DOCX parser (engine/src/application/parsers/docx_parser.py).

Uses python-docx to create real in-memory DOCX bytes — no mocking needed
because the parser itself has no external dependencies.
"""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from src.application.parsers.docx_parser import parse_docx
from src.domain.entities import ParsedDocument, DocumentSection


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_docx_bytes(paragraphs: list[tuple[str, str]]) -> bytes:
    """Create an in-memory DOCX from a list of (text, style_name) tuples."""
    doc = DocxDocument()
    for text, style in paragraphs:
        if style == "Normal":
            doc.add_paragraph(text)
        elif style == "Heading 1":
            doc.add_heading(text, level=1)
        elif style == "Heading 2":
            doc.add_heading(text, level=2)
        elif style == "Heading 3":
            doc.add_heading(text, level=3)
        elif style == "Heading 4":
            doc.add_heading(text, level=4)
        else:
            p = doc.add_paragraph(text)
            try:
                p.style = doc.styles[style]
            except KeyError:
                pass  # Style not available — paragraph remains Normal

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestParseDocxBasic:
    def test_returns_parsed_document(self):
        data = _make_docx_bytes(
            [
                ("Introduccion", "Heading 1"),
                ("Este es el cuerpo.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert isinstance(result, ParsedDocument)

    def test_full_text_contains_heading_and_body(self):
        data = _make_docx_bytes(
            [
                ("Introduccion", "Heading 1"),
                ("Este es el cuerpo.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert "Introduccion" in result.full_text
        assert "Este es el cuerpo." in result.full_text

    def test_markdown_heading_prefix(self):
        data = _make_docx_bytes(
            [
                ("Capitulo 1", "Heading 1"),
                ("Texto del capitulo.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert "# Capitulo 1" in result.markdown

    def test_heading2_uses_double_hash(self):
        data = _make_docx_bytes(
            [
                ("Seccion 1.1", "Heading 2"),
                ("Contenido.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert "## Seccion 1.1" in result.markdown

    def test_sections_list_populated(self):
        data = _make_docx_bytes(
            [
                ("Marco Teorico", "Heading 1"),
                ("Contenido del marco.", "Normal"),
                ("Subseccion", "Heading 2"),
                ("Detalles.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert len(result.sections) == 2

    def test_section_heading_and_level(self):
        data = _make_docx_bytes(
            [
                ("Marco Teorico", "Heading 1"),
                ("Contenido.", "Normal"),
            ]
        )
        result = parse_docx(data)
        section = result.sections[0]
        assert isinstance(section, DocumentSection)
        assert section.heading == "Marco Teorico"
        assert section.level == 1

    def test_section_content_trimmed(self):
        data = _make_docx_bytes(
            [
                ("Metodologia", "Heading 1"),
                ("Parrafo uno.", "Normal"),
                ("Parrafo dos.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert "Parrafo uno." in result.sections[0].content
        assert "Parrafo dos." in result.sections[0].content

    def test_section_offsets_are_ints(self):
        data = _make_docx_bytes(
            [
                ("Introduccion", "Heading 1"),
                ("Texto.", "Normal"),
            ]
        )
        result = parse_docx(data)
        section = result.sections[0]
        assert isinstance(section.offset_start, int)
        assert isinstance(section.offset_end, int)

    def test_offset_end_greater_than_start(self):
        data = _make_docx_bytes(
            [
                ("Introduccion", "Heading 1"),
                ("Texto.", "Normal"),
            ]
        )
        result = parse_docx(data)
        section = result.sections[0]
        assert section.offset_end > section.offset_start


class TestParseDocxEdgeCases:
    def test_empty_docx_returns_empty_sections(self):
        data = _make_docx_bytes([])
        result = parse_docx(data)
        assert result.sections == []
        assert result.full_text == ""

    def test_body_only_no_headings(self):
        """Document with no headings produces no sections but has full_text."""
        data = _make_docx_bytes(
            [
                ("Solo texto plano.", "Normal"),
                ("Otro parrafo.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert result.sections == []
        assert "Solo texto plano." in result.full_text

    def test_multiple_sections_sequential_offsets(self):
        data = _make_docx_bytes(
            [
                ("Capitulo 1", "Heading 1"),
                ("Contenido 1.", "Normal"),
                ("Capitulo 2", "Heading 1"),
                ("Contenido 2.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert len(result.sections) == 2
        # Second section must start after first one ends
        assert result.sections[1].offset_start >= result.sections[0].offset_end

    def test_heading3_level_correct(self):
        data = _make_docx_bytes(
            [
                ("Subsub", "Heading 3"),
                ("Detalle.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert result.sections[0].level == 3

    def test_heading4_level_correct(self):
        data = _make_docx_bytes(
            [
                ("Deep", "Heading 4"),
                ("Contenido.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert result.sections[0].level == 4

    def test_unicode_text_preserved(self):
        data = _make_docx_bytes(
            [
                ("Introducción", "Heading 1"),
                ("Análisis de métodos.", "Normal"),
            ]
        )
        result = parse_docx(data)
        assert "Introducción" in result.full_text
        assert "Análisis de métodos." in result.full_text
