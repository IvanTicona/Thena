from docx import Document as DocxDocument

from src.domain.entities import DocumentSection, ParsedDocument

HEADING_STYLES = {
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Title": 1,
}


def parse_docx(file_bytes: bytes) -> ParsedDocument:
    """Parse a DOCX file into structured sections with character offsets."""
    import io

    doc = DocxDocument(io.BytesIO(file_bytes))

    full_text_parts: list[str] = []
    markdown_parts: list[str] = []
    sections: list[DocumentSection] = []

    current_heading: str | None = None
    current_level: int = 0
    current_content_parts: list[str] = []
    current_offset_start: int = 0
    char_offset = 0

    def _flush_section():
        nonlocal current_heading, current_content_parts, current_offset_start
        if current_heading is not None:
            content = "\n".join(current_content_parts).strip()
            sections.append(
                DocumentSection(
                    heading=current_heading,
                    level=current_level,
                    content=content,
                    offset_start=current_offset_start,
                    offset_end=char_offset,
                )
            )
            current_content_parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ""
        heading_level = HEADING_STYLES.get(style_name)

        if heading_level:
            _flush_section()
            current_heading = text
            current_level = heading_level
            current_offset_start = char_offset

            md_prefix = "#" * heading_level
            markdown_parts.append(f"{md_prefix} {text}")
        else:
            current_content_parts.append(text)
            markdown_parts.append(text)

        full_text_parts.append(text)
        char_offset += len(text) + 1  # +1 for newline

    # Flush last section
    _flush_section()

    full_text = "\n".join(full_text_parts)
    markdown = "\n\n".join(markdown_parts)

    return ParsedDocument(full_text=full_text, markdown=markdown, sections=sections)
